# -*- coding: utf-8 -*-
"""T01 基线测试文档生成：1 份 PDF + 1 份 DOCX（长江文化史实内容）"""
from pathlib import Path

out_dir = Path(r"D:\长江学论纲\ycki\deploy\testdata")
out_dir.mkdir(parents=True, exist_ok=True)

# ---------- PDF（reportlab + CID 中文字体）----------
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont

pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

pdf_text = [
    ("汉阳铁厂创办始末", 16),
    ("汉阳铁厂是中国近代第一家钢铁联合企业，由湖广总督张之洞于1890年（光绪十六年）奏准创办，厂址位于湖北汉阳龟山北麓。", 11),
    ("张之洞调任湖广总督后，将原拟设于广东的炼铁厂改建于湖北。1891年汉阳铁厂破土动工，1893年建成，1894年5月正式投产。", 11),
    ("汉阳铁厂投产时拥有炼铁高炉两座，其规模在当时的远东首屈一指，被西方视为中国觉醒的标志。", 11),
    ("1896年，因经费困难，汉阳铁厂改为官督商办，由盛宣怀接办。1908年，盛宣怀将汉阳铁厂、大冶铁矿和萍乡煤矿合并，成立汉冶萍煤铁厂矿公司，这是中国近代最早的钢铁煤联营企业。", 11),
    ("汉阳铁厂的创办是长江中游近代工业文化的重要开端，带动了武汉地区民族工业的发展，汉阳也因此成为中国近代工业重镇。", 11),
]
c = canvas.Canvas(str(out_dir / "汉阳铁厂简介.pdf"), pagesize=A4)
w, h = A4
y = h - 80
for text, size in pdf_text:
    from reportlab.pdfbase.pdfmetrics import stringWidth
    max_chars = int((w - 100) / stringWidth("汉", "STSong-Light", size))
    while text:
        seg, text = text[:max_chars], text[max_chars:]
        c.setFont("STSong-Light", size)
        c.drawString(50, y, seg)
        y -= size + 8
    y -= 10
c.save()
print("PDF OK:", out_dir / "汉阳铁厂简介.pdf")

# ---------- DOCX（python-docx）----------
import docx

docx_text = [
    ("汉口开埠与近代长江航运", 0),
    ("1861年（咸丰十一年）3月，根据《天津条约》，汉口正式开埠对外通商，成为长江中游第一个通商口岸，英、俄、法、德、日等国相继在汉口设立租界。", 1),
    ("1872年（同治十一年），轮船招商局在上海成立，这是中国近代第一家轮船航运企业。1873年，轮船招商局开辟上海至汉口航线，打破了外国轮船公司对长江航运的垄断。", 1),
    ("汉口开埠后，对外贸易迅速发展，到20世纪初，汉口出口贸易占全国总额的十分之一以上，被誉为“东方芝加哥”。", 1),
    ("1906年，京汉铁路全线通车，汉口成为水陆交通枢纽，商业辐射范围进一步扩大，长江与铁路的联运奠定了武汉近代商业文化的格局。", 1),
]
d = docx.Document()
for text, level in docx_text:
    d.add_heading(text, level=1) if level == 0 else d.add_paragraph(text)
d.save(str(out_dir / "汉口开埠与长江航运.docx"))
print("DOCX OK:", out_dir / "汉口开埠与长江航运.docx")
